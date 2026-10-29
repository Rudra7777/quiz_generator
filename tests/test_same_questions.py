"""Tests for building every set's question order and option order in Same Questions mode."""

import io
from collections import Counter
from types import SimpleNamespace

import pandas as pd
import pytest
from openpyxl import load_workbook

from excel_handler import FullQuestion, FullQuestionBank
from response_generator import (
    SET_COL, answer_column, generate_responses, to_bank_columns, to_positional_columns,
)
from same_questions import build_same_question_sets, fixed_slots, held_in_place


def _question(no, options=None, answer="A", difficulty="medium") -> FullQuestion:
    a, b, c, d = options or (f"right {no}", f"wrong {no}.1", f"wrong {no}.2", f"wrong {no}.3")
    return FullQuestion(
        question_no=no, question_id=f"Q{no}", question_text=f"Question {no}?",
        option_a=a, option_b=b, option_c=c, option_d=d,
        answer=answer, difficulty=difficulty,
    )


@pytest.fixture
def bank20() -> FullQuestionBank:
    """20 questions, answers mostly A; Q5, Q9 and Q12 have options that refer to others."""
    levels = ["hard", "medium", "easy", "medium"]
    questions = [_question(no, difficulty=levels[no % 4]) for no in range(1, 21)]
    questions[4] = _question(5, ("one", "two", "three", "All of the above"), "D", "hard")
    questions[8] = _question(9, ("one", "two", "None of these", "four"), "A")
    questions[11] = _question(12, ("one", "two", "three", "Both A and B"), "A")
    return FullQuestionBank(questions)


def test_fixed_slots_detects_reference_options(bank20):
    q = {q.question_no: q for q in bank20.get_all()}
    assert fixed_slots(q[5]) == frozenset({3})
    assert fixed_slots(q[9]) == frozenset({2})
    assert fixed_slots(q[12]) == frozenset({0, 1, 2, 3})
    assert fixed_slots(q[1]) == frozenset()
    assert held_in_place(bank20) == [5, 9, 12]


def test_every_set_has_every_question_once(bank20):
    sets = build_same_question_sets(bank20, 65, seed=1)
    assert list(sets) == [f"S-{i:02d}" for i in range(1, 66)]
    for positions in sets.values():
        assert sorted(p.question_no for p in positions) == list(range(1, 21))


def test_no_two_sets_are_identical(bank20):
    sets = build_same_question_sets(bank20, 65, seed=1)
    assert len({tuple(sets[s]) for s in sets}) == 65


def test_slots_are_a_permutation_and_answer_points_at_the_bank_answer(bank20):
    bank = {q.question_no: q for q in bank20.get_all()}
    for positions in build_same_question_sets(bank20, 10, seed=2).values():
        for p in positions:
            assert sorted(p.slots) == list("ABCD")
            assert p.slots["ABCD".index(p.answer)] == bank[p.question_no].answer


def test_held_options_stay_in_their_slot(bank20):
    for positions in build_same_question_sets(bank20, 30, seed=3).values():
        by_no = {p.question_no: p for p in positions}
        assert by_no[5].slots[3] == "D" and by_no[5].answer == "D"
        assert by_no[9].slots[2] == "C"
        assert by_no[12].slots == "ABCD"


def test_each_set_key_is_balanced(bank20):
    # 20 questions -> 5 of each letter; Q5 (D) and Q12 (fixed) are absorbed.
    for positions in build_same_question_sets(bank20, 65, seed=4).values():
        counts = Counter(p.answer for p in positions)
        assert max(counts.values()) - min(counts.values()) <= 1


def test_seed_makes_generation_repeatable(bank20):
    assert build_same_question_sets(bank20, 5, seed=9) == build_same_question_sets(bank20, 5, seed=9)


def test_more_sets_than_possible_orders_is_refused():
    bank3 = FullQuestionBank([_question(no) for no in (1, 2, 3)])
    with pytest.raises(ValueError, match="only 6 different orders"):
        build_same_question_sets(bank3, 7, seed=0)


# ── Papers file ──────────────────────────────────────────────────────────────

from excel_export import (
    _load_question_bank_from_question_papers, create_formatted_excel,
    create_same_questions_excel,
)
from same_questions import read_set_map, variant_papers


def _papers_file(bank, tmp_path, num_sets=4, seed=1):
    sets = build_same_question_sets(bank, num_sets, seed=seed)
    path = tmp_path / "papers.xlsx"
    path.write_bytes(create_same_questions_excel(sets, bank, seed=seed))
    return sets, path


def test_same_questions_file_has_the_expected_sheets(bank20):
    data = create_same_questions_excel(build_same_question_sets(bank20, 4, seed=1), bank20, seed=1)
    wb = load_workbook(io.BytesIO(data))
    assert wb.sheetnames == ["All_Sets", "S-01", "S-02", "S-03", "S-04",
                             "Answer_Key", "Set_Map", "Question_Bank"]


def test_set_map_round_trips(bank20, tmp_path):
    sets, path = _papers_file(bank20, tmp_path)
    assert read_set_map(str(path)) == sets


def test_old_papers_have_no_set_map(tmp_path):
    bank = FullQuestionBank([_question(no, difficulty="hard") for no in range(1, 5)])
    matrix = [["Q1", "Q2"], ["Q3", "Q4"]]
    path = tmp_path / "old.xlsx"
    path.write_bytes(create_formatted_excel(matrix, matrix, {}, bank))
    assert read_set_map(str(path)) is None


def test_paper_prints_positions_and_shuffled_options_without_bank_numbers(bank20, tmp_path):
    sets, path = _papers_file(bank20, tmp_path)
    ws = load_workbook(path)["S-02"]
    bank = {q.question_no: q for q in bank20.get_all()}
    assert ws["C2"].value is None
    for k in range(1, 21):
        row = k + 2
        p = sets["S-02"][k - 1]
        assert ws.cell(row=row, column=1).value == k
        assert ws.cell(row=row, column=3).value is None
        assert ws.cell(row=row, column=2).value == p.question_no
        assert ws.cell(row=row, column=4).value == bank[p.question_no].question_text
        original = getattr(bank[p.question_no], f"option_{p.slots[0].lower()}")
        assert ws.cell(row=row, column=5).value == f"A] {original}"


def test_answer_key_is_per_set_by_position(bank20, tmp_path):
    sets, path = _papers_file(bank20, tmp_path)
    ws = load_workbook(path)["Answer_Key"]
    row = next(r for r in range(1, 12) if ws.cell(row=r, column=1).value == "S-03")
    for k in range(1, 21):
        assert ws.cell(row=row, column=k + 1).value == sets["S-03"][k - 1].answer


def test_question_bank_sheet_holds_the_bank_as_uploaded(bank20, tmp_path):
    _, path = _papers_file(bank20, tmp_path)
    loaded = _load_question_bank_from_question_papers(str(path))
    for got, want in zip(loaded.get_all(), bank20.get_all()):
        assert (got.question_no, got.question_text, got.option_a, got.option_b,
                got.option_c, got.option_d, got.answer) == (
            want.question_no, want.question_text, want.option_a, want.option_b,
            want.option_c, want.option_d, want.answer)
    assert len(loaded.get_all()) == 20


def test_variant_papers_print_options_in_set_order(bank20):
    sets = build_same_question_sets(bank20, 3, seed=5)
    variants, matrix = variant_papers(bank20, sets)
    bank = {q.question_no: q for q in bank20.get_all()}
    assert [len(ids) for ids in matrix] == [20, 20, 20]
    p = sets["S-02"][0]
    v = variants.get_by_id(matrix[1][0])
    assert v.question_id == f"S-02:{bank[p.question_no].question_id}"
    assert v.answer == p.answer and v.difficulty == bank[p.question_no].difficulty
    assert v.option_a == getattr(bank[p.question_no], f"option_{p.slots[0].lower()}")


# ── Translating answers between positions and bank questions ─────────────────



@pytest.fixture
def same_papers(bank20, tmp_path):
    sets, path = _papers_file(bank20, tmp_path, num_sets=6, seed=3)
    return SimpleNamespace(path=path, bank=bank20, sets=sets)


def _one_student(set_name, answers=None):
    row = {SET_COL: set_name}
    row.update({answer_column(i): None for i in range(1, 21)})
    row.update({answer_column(pos): letter for pos, letter in (answers or {}).items()})
    return pd.DataFrame([row])


def test_translation_round_trips(same_papers):
    df = generate_responses(str(same_papers.path), same_papers.bank, 6, seed=5)
    assert [c for c in df.columns if c.startswith("Q - ")] == [answer_column(i) for i in range(1, 21)]
    back = to_positional_columns(to_bank_columns(df, same_papers.sets), same_papers.sets)
    pd.testing.assert_frame_equal(back, df)


def test_typed_letter_maps_to_the_original_letter(same_papers):
    first = same_papers.sets["S-01"][0]
    bank_answer = next(q for q in same_papers.bank.get_all()
                       if q.question_no == first.question_no).answer
    out = to_bank_columns(_one_student("S-01", {1: first.answer}), same_papers.sets)
    assert out.loc[0, answer_column(first.question_no)] == bank_answer


def test_lowercase_and_spaces_are_read(same_papers):
    plain = to_bank_columns(_one_student("S-01", {1: "B"}), same_papers.sets)
    messy = to_bank_columns(_one_student("S-01", {1: " b"}), same_papers.sets)
    no = same_papers.sets["S-01"][0].question_no
    assert messy.loc[0, answer_column(no)] == plain.loc[0, answer_column(no)]


def test_blank_stays_blank(same_papers):
    out = to_bank_columns(_one_student("S-01"), same_papers.sets)
    assert out[[answer_column(i) for i in range(1, 21)]].isna().all().all()
    assert to_bank_columns(_one_student("S-01", {1: "E"}), same_papers.sets) \
        .loc[0, answer_column(same_papers.sets["S-01"][0].question_no)] == "E"


def test_unknown_set_is_refused(same_papers):
    with pytest.raises(ValueError, match="Unknown or missing Question Set"):
        to_bank_columns(_one_student("S-99"), same_papers.sets)


def test_extra_blank_columns_are_ignored_and_filled_ones_refused(same_papers):
    df = _one_student("S-01", {1: "A"})
    df["Q - 21 [Answer]"] = None
    out = to_bank_columns(df, same_papers.sets)
    assert "Q - 21 [Answer]" not in out.columns
    df.loc[0, "Q - 21 [Answer]"] = "A"
    with pytest.raises(ValueError, match=r"Q - 21 \[Answer\]"):
        to_bank_columns(df, same_papers.sets)


def test_non_contiguous_bank_numbers_round_trip(tmp_path):
    numbers = [3, 7, 15, 102, 4, 9, 11, 20, 33, 41, 58, 60, 71, 85, 90, 99, 105, 110, 111, 120]
    bank = FullQuestionBank([_question(no, difficulty="medium") for no in numbers])
    sets, path = _papers_file(bank, tmp_path, num_sets=6, seed=4)
    df = generate_responses(str(path), bank, 6, seed=5)
    assert [c for c in df.columns if c.startswith("Q - ")] == [answer_column(i) for i in range(1, 21)]
    bank_df = to_bank_columns(df, sets)
    assert [c for c in bank_df.columns if c.startswith("Q - ")] == [answer_column(n) for n in sorted(numbers)]
    pd.testing.assert_frame_equal(to_positional_columns(bank_df, sets), df)
    # every answer a student gave is a bank letter, and the generated correct rate maps back to the bank key
    assert bank_df[answer_column(3)].isin(list("ABCD")).all()


def test_positional_direction_refuses_filled_non_bank_column(tmp_path):
    bank = FullQuestionBank([_question(no) for no in range(101, 121)])
    sets, path = _papers_file(bank, tmp_path, num_sets=3, seed=1)
    df = generate_responses(str(path), bank, 3, seed=1)
    bank_df = to_bank_columns(df, sets)
    bank_df["Q - 05 [Answer]"] = None
    assert "Q - 05 [Answer]" not in to_positional_columns(bank_df, sets).columns.difference(
        [answer_column(i) for i in range(1, 21)])
    bank_df.loc[0, "Q - 05 [Answer]"] = "A"
    with pytest.raises(ValueError, match="not a question in the bank"):
        to_positional_columns(bank_df, sets)


# ── Scoring Same Questions responses ─────────────────────────────────────────

from answer_checker import (
    check_all_responses, generate_scoring_report, normalize_response_sheet,
)
from response_generator import ROLL_COL


def _score(df, path, bank):
    return check_all_responses(normalize_response_sheet(df), str(path), bank)


def _scored_workbook(report, path, bank, tmp_path):
    out = tmp_path / "report.xlsx"
    generate_scoring_report(report, str(out), str(path), bank)
    return load_workbook(out)


def test_scores_match_what_the_generator_intended(same_papers):
    df = generate_responses(str(same_papers.path), same_papers.bank, 6, correct_rate=0.6, seed=7)
    report = _score(df, same_papers.path, same_papers.bank)
    by_roll = {r.roll_number: r for r in report.student_reports}
    assert len(by_roll) == len(df)
    for _, row in df.iterrows():
        key = [p.answer for p in same_papers.sets[row["Question Set"]]]
        typed = [row[answer_column(i)] for i in range(1, 21)]
        r = by_roll[str(row[ROLL_COL]).strip()]
        assert r.correct == sum(t == k for t, k in zip(typed, key))
        assert r.assigned == 20 and r.validation.extra_count == 0


def test_blank_position_is_wrong(same_papers):
    key = [p.answer for p in same_papers.sets["S-01"]]
    df = _one_student("S-01", {i: key[i - 1] for i in range(1, 20)})  # position 20 left blank
    df[ROLL_COL] = "R1"
    df["Timestamp"] = pd.Timestamp("2026-01-01")
    report = _score(df, same_papers.path, same_papers.bank)
    r = report.student_reports[0]
    assert (r.correct, r.wrong, r.assigned) == (19, 1, 20)


def test_faculty_report_is_by_bank_question_with_the_bank_key(same_papers, tmp_path):
    df = generate_responses(str(same_papers.path), same_papers.bank, 6, seed=7)
    report = _score(df, same_papers.path, same_papers.bank)
    ws = _scored_workbook(report, same_papers.path, same_papers.bank, tmp_path)["Faculty_Report"]
    strip = [c.value for c in ws[1] if c.value]
    key = [c.value for c in ws[2] if c.value]
    assert strip == [f"Q{n:02d}" for n in range(1, 21)]
    assert key == [q.answer for q in sorted(same_papers.bank.get_all(), key=lambda q: q.question_no)]
    # no extra-answer (yellow) rule fires: every student answered only their own questions
    assert all(r.validation.extra_count == 0 for r in report.student_reports)


def test_as_typed_sheet(same_papers, tmp_path):
    df = generate_responses(str(same_papers.path), same_papers.bank, 6, correct_rate=0.5, seed=7)
    report = _score(df, same_papers.path, same_papers.bank)
    wb = _scored_workbook(report, same_papers.path, same_papers.bank, tmp_path)
    assert wb.sheetnames.index("As_Typed") == wb.sheetnames.index("Faculty_Report") + 1
    ws = wb["As_Typed"]
    assert ws["A1"].value.startswith("Coloured against each set's key")
    assert [c.value for c in ws[3]] == (
        ["Roll Number", "Full Name", "Set"] + [f"Q{i}" for i in range(1, 21)] + ["Correct"])
    scores = {str(row[0].value): row[6].value for row in wb["Scores"].iter_rows(min_row=2)}
    typed_by_roll = {str(row[ROLL_COL]).strip(): row for _, row in df.iterrows()}
    rows = list(ws.iter_rows(min_row=4, values_only=False))
    assert len(rows) == len(df)
    for cells in rows:
        roll = str(cells[0].value)
        src = typed_by_roll[roll]
        key = [p.answer for p in same_papers.sets[cells[2].value]]
        for i in range(1, 21):
            cell = cells[2 + i]
            assert cell.value == src[answer_column(i)]
            if cell.value is not None:
                assert cell.font.italic == (cell.value != key[i - 1])
        assert cells[23].value == scores[roll] and isinstance(cells[23].value, int)


def test_as_typed_leaves_blank_unfilled(same_papers, tmp_path):
    df = _one_student("S-01", {1: "A"})
    df[ROLL_COL] = "R1"
    df["Timestamp"] = pd.Timestamp("2026-01-01")
    report = _score(df, same_papers.path, same_papers.bank)
    ws = _scored_workbook(report, same_papers.path, same_papers.bank, tmp_path)["As_Typed"]
    assert ws.cell(row=4, column=5).value is None
    assert ws.cell(row=4, column=5).fill.fill_type is None


def test_non_contiguous_bank_scores_end_to_end(tmp_path):
    numbers = [3, 7, 15, 102, 4, 9, 11, 20, 33, 41, 58, 60, 71, 85, 90, 99, 105, 110, 111, 120]
    bank = FullQuestionBank([_question(no) for no in numbers])
    sets, path = _papers_file(bank, tmp_path, num_sets=4, seed=2)
    df = generate_responses(str(path), bank, 4, correct_rate=0.7, seed=3)
    report = _score(df, path, bank)
    by_roll = {r.roll_number: r for r in report.student_reports}
    for _, row in df.iterrows():
        key = [p.answer for p in sets[row["Question Set"]]]
        typed = [row[answer_column(i)] for i in range(1, 21)]
        r = by_roll[str(row[ROLL_COL]).strip()]
        assert r.correct == sum(t == k for t, k in zip(typed, key))
        assert r.validation.extra_count == 0
    wb = _scored_workbook(report, path, bank, tmp_path)
    strip = [c.value for c in wb["Faculty_Report"][1] if c.value]
    assert strip == [f"Q{n:02d}" for n in sorted(numbers)]
    assert wb["As_Typed"].max_row == 3 + len(df)


def test_old_papers_have_no_as_typed_sheet(tmp_path):
    from pathlib import Path
    from excel_handler import load_question_bank
    root = Path(__file__).resolve().parents[1]
    bank = load_question_bank(str(root / "input" / "question_bank_72.xlsx"))
    papers = root / "output" / "question_papers.xlsx"
    df = generate_responses(str(papers), bank, 5, seed=1)
    report = check_all_responses(normalize_response_sheet(df), str(papers), bank)
    out = tmp_path / "report.xlsx"
    generate_scoring_report(report, str(out), str(papers), bank)
    wb = load_workbook(out)
    assert "Faculty_Report" in wb.sheetnames and "As_Typed" not in wb.sheetnames


@pytest.mark.parametrize("text", [
    "Both A and B", "A and C", "Only B", "A & B", "A, B and C", "Neither A nor B",
    "A or B", "Both A and B are correct", "Only C is correct",
])
def test_options_naming_letters_hold_every_slot(text):
    q = _question(1, ("one", "two", "three", text))
    assert fixed_slots(q) == frozenset({0, 1, 2, 3})


@pytest.mark.parametrize("text", [
    "A capability a firm performs distinctively well and that rivals struggle to imitate",
    "The network of organisations involved in producing and delivering a product",
    "A queue of customers waiting for a service",
    "A dog and a cat", "Vitamin A and C",
])
def test_ordinary_sentences_with_letter_words_are_not_held(text):
    assert fixed_slots(_question(1, ("one", "two", "three", text))) == frozenset()


@pytest.mark.parametrize("text", ["all of these", "neither of the above"])
def test_reference_phrases_hold_only_their_slot(text):
    assert fixed_slots(_question(1, ("one", "two", "three", text))) == frozenset({3})


def test_real_bank_first_20_questions_are_not_held(tmp_path):
    from excel_handler import load_question_bank
    path = tmp_path / "bank20.xlsx"
    pd.read_excel("input/question_bank.xlsx").head(20).to_excel(path, index=False)
    assert held_in_place(load_question_bank(str(path))) == []


# --- Set_Map robustness ------------------------------------------------------

def _edit_set_map(path, edit):
    wb = load_workbook(path)
    ws = wb["Set_Map"]
    edit(ws)
    wb.save(path)


def _data_rows(ws):
    return [[c.value for c in r] for r in ws.iter_rows(min_row=4, max_row=ws.max_row)
            if r[0].value is not None]


def test_set_map_survives_sorting_the_sheet(bank20, tmp_path):
    sets, path = _papers_file(bank20, tmp_path)

    def resort(ws):
        rows = sorted(_data_rows(ws), key=lambda r: (r[2], r[0]))
        for i, values in enumerate(rows, 4):
            for col, v in enumerate(values, 1):
                ws.cell(row=i, column=col, value=v)

    _edit_set_map(path, resort)
    assert read_set_map(str(path)) == sets


def test_broken_position_column_raises(bank20, tmp_path):
    _, path = _papers_file(bank20, tmp_path)
    _edit_set_map(path, lambda ws: ws.cell(row=5, column=2, value=1))   # duplicate position
    with pytest.raises(ValueError, match="positions"):
        read_set_map(str(path))


def test_sets_of_different_lengths_raise(bank20, tmp_path):
    _, path = _papers_file(bank20, tmp_path)
    _edit_set_map(path, lambda ws: ws.delete_rows(ws.max_row))   # drop S-04's last row
    with pytest.raises(ValueError, match="same number"):
        read_set_map(str(path))


def test_note_row_is_ignored_but_bad_data_row_is_named(bank20, tmp_path):
    sets, path = _papers_file(bank20, tmp_path)
    _edit_set_map(path, lambda ws: ws.cell(row=ws.max_row + 2, column=1, value="remember to check"))
    assert read_set_map(str(path)) == sets

    _edit_set_map(path, lambda ws: ws.cell(row=6, column=3, value="oops"))
    with pytest.raises(ValueError, match="row 6"):
        read_set_map(str(path))

    _edit_set_map(path, lambda ws: (ws.cell(row=6, column=3, value=1),
                                    setattr(ws.cell(row=6, column=8), "value", None)))
    with pytest.raises(ValueError, match="row 6"):
        read_set_map(str(path))
