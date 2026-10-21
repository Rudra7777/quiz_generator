"""
Option Shuffler

Reorders each question's A-D options so that the correct answer lands on every
letter equally often — 25% each — within each difficulty level. A bank whose key
leans on one letter lets students score by guessing that letter.

Works on the bank as a DataFrame in the question bank file format, so the result
can be saved and fed straight back into Generate Papers.
"""

import io
import random
from itertools import permutations
from typing import Optional

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

from excel_handler import normalize_difficulty


LETTERS = "ABCD"
OPTION_COLS = ["option_a", "option_b", "option_c", "option_d"]

MOVED_FILL = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
STAYED_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
THIN = Side(style="thin")


def _column_lookup(df: pd.DataFrame) -> dict:
    """Map normalized column names (as load_question_bank reads them) to the real ones."""
    return {str(c).strip().lower().replace(" ", "_"): c for c in df.columns}


def answer_counts(df: pd.DataFrame) -> pd.DataFrame:
    """Correct-answer counts: one row per difficulty (Hard/Medium/Easy), one column per letter."""
    cols = _column_lookup(df)
    difficulty = df[cols["difficulty"]].map(lambda v: normalize_difficulty(v).capitalize())
    answer = df[cols["answer"]].astype(str).str.strip().str.upper()
    counts = pd.crosstab(difficulty, answer).reindex(
        index=["Hard", "Medium", "Easy"], columns=list(LETTERS), fill_value=0
    )
    return counts.loc[counts.sum(axis=1) > 0]


def _avoid_current_letter(targets: list, current: list, rng: random.Random) -> None:
    """
    Swap targets in place until no question keeps its current answer letter.

    A swap with a question whose current and target letters are both different from
    the clash always exists unless one letter is more than 75% of the level's key;
    in that case the clashes that cannot be removed are left alone.
    """
    for i in range(len(targets)):
        if targets[i] != current[i]:
            continue
        swaps = [
            j for j in range(len(targets))
            if targets[j] != current[i] and current[j] != targets[i]
        ]
        if swaps:
            j = rng.choice(swaps)
            targets[i], targets[j] = targets[j], targets[i]


def _new_order(current: int, target: int, rng: random.Random) -> tuple:
    """
    Pick a new order for the four options: new_order[slot] is the old slot whose
    option goes there. The correct option goes from current to target and no other
    option stays where it was.
    """
    orders = [
        order for order in permutations(range(4))
        if order[target] == current
        and all(order[slot] != slot for slot in range(4) if slot != target)
    ]
    return rng.choice(orders)


def balance_options(df: pd.DataFrame, seed: Optional[int] = None) -> pd.DataFrame:
    """
    Return a copy of the bank with options reordered so the correct answers are
    spread evenly over A-D within each difficulty.

    Each difficulty gets floor(n/4) of every letter; its leftover 1-3 questions go to
    whichever letters the bank has fewest of so far, so the bank as a whole comes out
    even too. No question keeps its old answer letter and no distractor stays in its
    old place, unless one letter makes up more than 75% of a level's key — then a
    few answers have to stay put. Question text, numbers and difficulty are untouched.
    """
    cols = _column_lookup(df)
    missing = [c for c in OPTION_COLS + ["answer", "difficulty"] if c not in cols]
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    rng = random.Random(seed)
    out = df.copy()
    option_cols = [cols[c] for c in OPTION_COLS]
    answer_col = cols["answer"]
    difficulty = df[cols["difficulty"]].map(normalize_difficulty)

    totals = {letter: 0 for letter in LETTERS}
    for level in ("hard", "medium", "easy"):
        rows = list(df.index[difficulty == level])
        if not rows:
            continue

        base, extra = divmod(len(rows), 4)
        # Leftovers go to the letters the bank is shortest on; ties broken at random.
        shortest = sorted(LETTERS, key=lambda letter: (totals[letter], rng.random()))
        targets = []
        for letter in LETTERS:
            share = base + (1 if letter in shortest[:extra] else 0)
            targets += [letter] * share
            totals[letter] += share
        rng.shuffle(targets)

        current = []
        for row in rows:
            letter = str(df.at[row, answer_col]).strip().upper()
            if letter not in LETTERS:
                raise ValueError(
                    f"Question at row {row + 2} has answer '{letter}', expected A/B/C/D."
                )
            current.append(letter)
        _avoid_current_letter(targets, current, rng)

        for row, letter, target in zip(rows, current, targets):
            options = [df.at[row, c] for c in option_cols]
            order = _new_order(LETTERS.index(letter), LETTERS.index(target), rng)

            for col, old_slot in zip(option_cols, order):
                out.at[row, col] = options[old_slot]
            out.at[row, answer_col] = target

    return out


def shuffle_check(original: pd.DataFrame, balanced: pd.DataFrame) -> dict:
    """
    Compare a bank with its balanced copy, question by question: how many changed
    their answer letter, how many have all four options in new places, and the
    question numbers of any that did not fully move.
    """
    cols = _column_lookup(original)
    option_cols = [cols[c] for c in OPTION_COLS]
    answer_col = cols["answer"]

    same_answer = (
        original[answer_col].astype(str).str.strip().str.upper()
        == balanced[answer_col].astype(str).str.strip().str.upper()
    )
    unmoved_option = (original[option_cols].values == balanced[option_cols].values).any(axis=1)
    not_moved = original.index[unmoved_option | same_answer.values]

    numbers = original[cols["question_no"]]
    return {
        "total": len(original),
        "answer_moved": int((~same_answer).sum()),
        "all_moved": int(len(original) - len(not_moved)),
        "not_moved": [int(n) for n in numbers.loc[not_moved]],
    }


def balanced_bank_bytes(original: pd.DataFrame, balanced: pd.DataFrame) -> bytes:
    """
    The balanced bank as an Excel file with three sheets:

        Shuffled     the balanced bank — first, because Generate Papers reads the
                     first sheet, so this file can be uploaded there as it is
        Original     the bank as uploaded
        Comparison   each question's old options over its new ones, so every
                     option can be checked by eye
    """
    output = io.BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        balanced.to_excel(writer, sheet_name="Shuffled", index=False)
        original.to_excel(writer, sheet_name="Original", index=False)
        _write_comparison(writer.book.create_sheet("Comparison"), original, balanced)
    return output.getvalue()


def _write_comparison(ws, original: pd.DataFrame, balanced: pd.DataFrame) -> None:
    """
    Two rows per question: the old options, then the new ones under them. A new
    option is green where it differs from the old one in that column and red where
    it is the same; the correct option is bold in both rows. Moved says YES only when
    the answer letter changed and all four options are in new places.
    """
    cols = _column_lookup(original)
    option_cols = [cols[c] for c in OPTION_COLS]
    answer_col = cols["answer"]

    headers = ["Q.No", "Question", "Row", "A", "B", "C", "D", "Ans", "Moved"]
    for col, header in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.font = Font(name="Calibri", bold=True)
        cell.alignment = Alignment(horizontal="center")

    wrap = Alignment(wrap_text=True, vertical="top")
    center = Alignment(horizontal="center", vertical="top")
    for i, (idx, old) in enumerate(original.iterrows()):
        new = balanced.loc[idx]
        old_row, new_row = 2 + 2 * i, 3 + 2 * i
        old_answer = str(old[answer_col]).strip().upper()
        new_answer = str(new[answer_col]).strip().upper()
        same = [old[c] == new[c] for c in option_cols]
        moved = old_answer != new_answer and not any(same)

        ws.cell(old_row, 1, old[cols["question_no"]]).alignment = center
        ws.cell(old_row, 2, old[cols["question"]]).alignment = wrap
        ws.merge_cells(start_row=old_row, start_column=1, end_row=new_row, end_column=1)
        ws.merge_cells(start_row=old_row, start_column=2, end_row=new_row, end_column=2)

        for row, values, answer in ((old_row, old, old_answer), (new_row, new, new_answer)):
            ws.cell(row, 3, "Old" if row == old_row else "New").alignment = center
            for offset, col in enumerate(option_cols):
                cell = ws.cell(row, 4 + offset, values[col])
                cell.alignment = wrap
                cell.font = Font(name="Calibri", bold=LETTERS[offset] == answer)
                if row == new_row:
                    cell.fill = STAYED_FILL if same[offset] else MOVED_FILL
            ws.cell(row, 8, answer).alignment = center

        flag = ws.cell(new_row, 9, "YES" if moved else "NO")
        flag.alignment = center
        flag.font = Font(name="Calibri", bold=True)
        flag.fill = MOVED_FILL if moved else STAYED_FILL
        ws.cell(new_row, 8).fill = STAYED_FILL if old_answer == new_answer else MOVED_FILL

        # One box round each question's pair of rows.
        for row in (old_row, new_row):
            for col in range(1, len(headers) + 1):
                ws.cell(row, col).border = Border(
                    left=THIN, right=THIN,
                    top=THIN if row == old_row else None,
                    bottom=THIN if row == new_row else None,
                )

    for letter, width in zip("ABCDEFGHI", (6, 40, 6, 28, 28, 28, 28, 6, 8)):
        ws.column_dimensions[letter].width = width
    ws.freeze_panes = "D2"
